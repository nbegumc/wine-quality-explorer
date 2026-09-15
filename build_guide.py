"""Build the Word and PDF editions of the beginner guide from its Markdown source.

Usage: uv run --with python-docx --with reportlab build_guide.py
Reads docs/BEGINNER_GUIDE.md and writes app/wine-quality-learning-guide.docx, using
the styles stored in docs/guide-template.docx, and app/wine-quality-learning-guide.pdf.
The guide's Markdown dialect is small: a title line, headings, paragraphs, bullet and
numbered lists, pipe tables, figures, and inline bold, code, and links.
"""
import copy
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "BEGINNER_GUIDE.md"
TEMPLATE = ROOT / "docs" / "guide-template.docx"
TARGET = ROOT / "app" / "wine-quality-learning-guide.docx"
PDF_TARGET = ROOT / "app" / "wine-quality-learning-guide.pdf"
TEXT_WIDTH_TWIPS = 9994      # Letter page with the template's margins
FIGURE_WIDTH = Emu(6345936)  # the same width in EMU
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(paragraph, text):
    """Append runs for Markdown inline bold, code, and links."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif piece.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece).groups()
            add_hyperlink(paragraph, label, url)
        else:
            paragraph.add_run(piece)


def add_hyperlink(paragraph, label, url):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "752350")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def set_spacing(paragraph, **attrs):
    spacing = OxmlElement("w:spacing")
    for key, value in attrs.items():
        spacing.set(qn(f"w:{key}"), str(value))
    paragraph.paragraph_format.element.get_or_add_pPr().append(spacing)


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    header, aligns, body = rows[0], rows[1], rows[2:]
    numeric = re.compile(r"^[\d.,%<>= –-]+$|^Less than")
    weights = [min(40, max(len(header[i]), *(len(r[i]) for r in body))) + 4 for i in range(len(header))]
    widths = [round(TEXT_WIDTH_TWIPS * w / sum(weights)) for w in weights]
    widths[-1] += TEXT_WIDTH_TWIPS - sum(widths)
    table = document.add_table(rows=len(rows) - 1, cols=len(header))
    table.alignment = 1  # centered
    # Schema order inside tblPr: ... jc, tblBorders, tblLayout, tblLook.
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    look.addprevious(copy.deepcopy(TABLE_BORDERS))
    look.addprevious(layout)
    for grid, width in zip(table._tbl.findall(qn("w:tblGrid") + "/" + qn("w:gridCol")), widths):
        grid.set(qn("w:w"), str(width))
    for r, values in enumerate([header] + body):
        row = table.rows[r]
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if r == 0:
            row_props.append(OxmlElement("w:tblHeader"))
        fill = "273649" if r == 0 else ("FFFFFF" if r % 2 else "F3F5F8")
        for c, value in enumerate(values):
            cell = row.cells[c]
            cell.width = Twips(widths[c])
            # Schema order inside tcPr: tcW, shd, tcMar, vAlign.
            cell_props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:fill"), fill)
            cell_props.append(shade)
            cell_props.append(copy.deepcopy(CELL_MARGINS))
            valign = OxmlElement("w:vAlign")
            valign.set(qn("w:val"), "center")
            cell_props.append(valign)
            paragraph = cell.paragraphs[0]
            if r == 0:
                paragraph.paragraph_format.keep_with_next = True
            set_spacing(paragraph, after=0, line=245, lineRule="auto")
            centered = r > 0 and (aligns[c].endswith(":") or numeric.match(value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
            add_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                if r == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    spacer = document.add_paragraph()
    set_spacing(spacer, after=40, before=0, line=240, lineRule="auto")
    spacer.add_run().font.size = Pt(2)


def build_borders():
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    return borders


def build_margins():
    margins = OxmlElement("w:tcMar")
    for side, width in (("top", 65), ("left", 85), ("bottom", 65), ("right", 85)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(width))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    return margins


TABLE_BORDERS = build_borders()
CELL_MARGINS = build_margins()


def parse(lines):
    """Yield (kind, payload) blocks from the guide's Markdown."""
    title_seen = subtitle_seen = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            title_seen = True
            yield "title", line[2:].strip()
        elif line.startswith("## "):
            yield "h1", line[3:].strip()
        elif line.startswith("### "):
            yield "h2", line[4:].strip()
        elif line.startswith("- "):
            yield "bullet", line[2:].strip()
        elif re.match(r"^\d+\. ", line):
            yield "numbered", line.strip()
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            yield "table", block
            continue
        elif line.startswith("!["):
            alt, path = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line).groups()
            yield "figure", (alt, SOURCE.parent / path)
        elif title_seen and not subtitle_seen:
            subtitle_seen = True
            yield "subtitle", line.strip()
        else:
            yield "paragraph", line.strip()
        i += 1


def build_docx(blocks):
    document = Document(str(TEMPLATE))
    title = subtitle = None
    for kind, payload in blocks:
        if kind == "title":
            title = payload
            document.add_paragraph(payload, style="Title")
        elif kind == "subtitle":
            subtitle = payload
            document.add_paragraph(payload, style="Subtitle")
        elif kind == "h1":
            document.add_paragraph(payload, style="Heading 1").paragraph_format.page_break_before = True
        elif kind == "h2":
            document.add_paragraph(payload, style="Heading 2")
        elif kind == "bullet":
            add_runs(document.add_paragraph(style="List Bullet"), payload)
        elif kind == "numbered":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Emu(288 * 635)
            paragraph.paragraph_format.first_line_indent = Emu(-288 * 635)
            add_runs(paragraph, payload)
        elif kind == "table":
            add_table(document, payload)
        elif kind == "figure":
            alt, path = payload
            document.add_picture(str(path), width=FIGURE_WIDTH)
            figure = document.paragraphs[-1]
            figure.paragraph_format.keep_with_next = True
            set_spacing(figure, after=120)
            figure._p.find(".//" + qn("wp:docPr")).set("descr", alt)
        else:
            add_runs(document.add_paragraph(), payload)
    core = document.core_properties
    core.title, core.subject, core.author = title, subtitle, "Wine Quality Explorer"
    core.keywords = "Bayesian network, probability, Python, wine quality, learning guide"
    core.description = "generated by python-docx from docs/BEGINNER_GUIDE.md"
    document.save(str(TARGET))
    print(f"Wrote {TARGET.name}: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables.")


def markup(text):
    """Convert the inline Markdown to reportlab paragraph markup."""
    out = []
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**"):
            out.append(f"<b>{html.escape(piece[2:-2])}</b>")
        elif piece.startswith("`"):
            out.append(f'<font face="Courier" size="9.5">{html.escape(piece[1:-1])}</font>')
        elif piece.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece).groups()
            out.append(f'<link href="{html.escape(url)}" color="#752350"><u>{html.escape(label)}</u></link>')
        else:
            out.append(html.escape(piece))
    return "".join(out)


def build_pdf(blocks):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=7)
    styles = {
        "title": ParagraphStyle("title", parent=body, fontName="Helvetica-Bold", fontSize=26, leading=31, spaceAfter=10),
        "subtitle": ParagraphStyle("subtitle", parent=body, fontSize=13, leading=17, textColor=colors.HexColor("#4B5B6D"), spaceAfter=16),
        "h1": ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=19, leading=24, spaceBefore=4, spaceAfter=9, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=9, spaceAfter=7, keepWithNext=True),
        "bullet": ParagraphStyle("bullet", parent=body, leftIndent=16, bulletIndent=4, spaceAfter=5),
        "numbered": ParagraphStyle("numbered", parent=body, leftIndent=16, firstLineIndent=-16),
        "cell": ParagraphStyle("cell", parent=body, fontSize=9.5, leading=12, spaceAfter=0),
        "head": ParagraphStyle("head", parent=body, fontName="Helvetica-Bold", fontSize=9.5, leading=12, spaceAfter=0, textColor=colors.white),
    }
    width = letter[0] - 2 * 0.78 * inch
    story, title, subtitle = [], None, None
    first_chapter = True
    for kind, payload in blocks:
        if kind == "title":
            title = payload
            story.append(Paragraph(markup(payload), styles["title"]))
        elif kind == "subtitle":
            subtitle = payload
            story.append(Paragraph(markup(payload), styles["subtitle"]))
        elif kind == "h1":
            story.append(PageBreak())
            story.append(Paragraph(markup(payload), styles["h1"]))
            first_chapter = False
        elif kind == "h2":
            story.append(Paragraph(markup(payload), styles["h2"]))
        elif kind == "bullet":
            story.append(Paragraph(markup(payload), styles["bullet"], bulletText="\u2022"))
        elif kind == "numbered":
            story.append(Paragraph(markup(payload), styles["numbered"]))
        elif kind == "table":
            rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in payload]
            header, aligns, data = rows[0], rows[1], rows[2:]
            numeric = re.compile(r"^[\d.,%<>= \u2013-]+$|^Less than")
            weights = [min(40, max(len(header[c]), *(len(r[c]) for r in data))) + 4 for c in range(len(header))]
            widths = [width * w / sum(weights) for w in weights]
            cells = [[Paragraph(markup(v), styles["head"]) for v in header]]
            for row in data:
                cells.append([Paragraph(markup(v), ParagraphStyle("c", parent=styles["cell"], alignment=TA_CENTER)
                                        if aligns[c].endswith(":") or numeric.match(v) else styles["cell"])
                              for c, v in enumerate(row)])
            table = Table(cells, colWidths=widths, repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#273649")),
                     ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                     ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                     ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
            style += [("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F3F5F8")) for r in range(2, len(cells), 2)]
            table.setStyle(TableStyle(style))
            story += [table, Spacer(1, 8)]
        elif kind == "figure":
            alt, path = payload
            image = Image(str(path))
            image.drawHeight = image.drawHeight * width / image.drawWidth
            image.drawWidth = width
            story.append(KeepTogether([image, Spacer(1, 6)]))
        else:
            story.append(Paragraph(markup(payload), body))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8.5)
        canvas.setFillColor(colors.HexColor("#617082"))
        canvas.drawString(0.78 * inch, 0.45 * inch, title or "")
        canvas.drawRightString(letter[0] - 0.78 * inch, 0.45 * inch, str(doc.page))
        canvas.restoreState()

    document = SimpleDocTemplate(str(PDF_TARGET), pagesize=letter, leftMargin=0.78 * inch, rightMargin=0.78 * inch,
                                 topMargin=0.7 * inch, bottomMargin=0.75 * inch, title=title, subject=subtitle,
                                 author="Wine Quality Explorer")
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    print(f"Wrote {PDF_TARGET.name}: {document.page} pages.")


def main():
    blocks = list(parse(SOURCE.read_text(encoding="utf-8").splitlines()))
    build_docx(blocks)
    build_pdf(blocks)


if __name__ == "__main__":
    main()
